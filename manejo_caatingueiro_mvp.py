"""
Manejo Caatingueiro - Ferramenta de Coleta (Concierge MVP)
============================================================

Este script NAO e o aplicativo completo do Manejo Caatingueiro. Ele existe
para viabilizar o "Concierge MVP" descrito na estrategia de lancamento:

  - Um agente de campo (ou voce mesmo) registra, manualmente, o nivel do
    reservatorio (poco/cisterna) relatado semanalmente por 5 a 10
    familias-piloto (via WhatsApp ou visita presencial).
  - O mesmo agente registra a chuva acumulada do mes por localidade,
    cruzando com dado publico (INMET/INPE) consultado manualmente.
  - O sistema gera uma SUGESTAO consultiva de plantio (nunca uma
    determinacao) comparando a chuva acumulada com o limiar da cultura
    da familia - sempre mostrando o dado que fundamenta a sugestao,
    respeitando a regra de negocio RN12 da especificacao de requisitos:
    a decisao final e sempre do agricultor.

Nao ha canal de solicitacao de agua aqui de proposito: essa funcionalidade
esta no Escopo Futuro / Pos-MVP e nao deve gerar expectativa neste piloto.

Todo o "banco de dados" e um unico arquivo Excel (.xlsx), criado
automaticamente na primeira execucao, com uma aba por entidade. Isso
permite abrir e inspecionar os dados a qualquer momento no Excel/LibreOffice,
sem precisar de nenhum servidor ou banco de dados de verdade - adequado ao
volume de 5 a 10 familias do piloto.

Requisitos: Python 3.8+ e a biblioteca "openpyxl" (pip install openpyxl).
Uso: python manejo_caatingueiro_mvp.py
"""

from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:
    print("A biblioteca 'openpyxl' nao esta instalada.")
    print("Instale com: pip install openpyxl")
    sys.exit(1)


ARQUIVO_DADOS = Path(__file__).resolve().parent / "dados_manejo_caatingueiro.xlsx"

# Cabecalhos de cada aba da planilha (a ordem define a ordem das colunas)
ABAS = {
    "Familias": ["id", "nome", "localidade", "cultura_principal", "contato"],
    "Localidades": ["nome", "media_historica_mm"],
    "Culturas": ["nome", "limiar_chuva_mm"],
    "RegistroChuva": ["id", "localidade", "mes_referencia", "volume_mm", "data_registro"],
    "RegistroReservatorio": ["id", "familia_id", "data", "nivel_percentual", "observacao"],
    "SugestoesPlantio": [
        "id", "familia_id", "data", "cultura", "chuva_acumulada_mm",
        "media_historica_mm", "percentual_da_media", "mensagem",
    ],
}


# ---------------------------------------------------------------------------
# Camada de acesso ao "banco de dados" (o arquivo Excel)
# ---------------------------------------------------------------------------

def inicializar_planilha() -> None:
    """Cria o arquivo Excel com todas as abas e cabecalhos, se ainda nao existir."""
    if ARQUIVO_DADOS.exists():
        return

    wb = Workbook()
    primeira_aba = True
    for nome_aba, colunas in ABAS.items():
        if primeira_aba:
            ws = wb.active
            ws.title = nome_aba
            primeira_aba = False
        else:
            ws = wb.create_sheet(nome_aba)
        ws.append(colunas)
    wb.save(ARQUIVO_DADOS)
    print(f"Planilha criada em: {ARQUIVO_DADOS}")


def abrir_planilha() -> Workbook:
    inicializar_planilha()
    return load_workbook(ARQUIVO_DADOS)


def salvar_planilha(wb: Workbook) -> None:
    wb.save(ARQUIVO_DADOS)


def proximo_id(ws: Worksheet) -> int:
    """Calcula o proximo id sequencial de uma aba que tem 'id' na primeira coluna."""
    maior = 0
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha[0] is not None:
            maior = max(maior, int(linha[0]))
    return maior + 1


def linhas_como_dicts(ws: Worksheet) -> list[dict]:
    """Converte as linhas de uma aba em uma lista de dicionarios {coluna: valor}."""
    cabecalho = [c.value for c in ws[1]]
    resultado = []
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in linha):
            continue
        resultado.append(dict(zip(cabecalho, linha)))
    return resultado


# ---------------------------------------------------------------------------
# Entrada de dados auxiliares (com validacao simples)
# ---------------------------------------------------------------------------

def pedir_texto(pergunta: str) -> str:
    while True:
        valor = input(pergunta).strip()
        if valor:
            return valor
        print("  -> Este campo nao pode ficar vazio.")


def pedir_numero(pergunta: str) -> float:
    while True:
        bruto = input(pergunta).strip().replace(",", ".")
        try:
            return float(bruto)
        except ValueError:
            print("  -> Digite um numero valido (ex.: 42.5).")


def pedir_inteiro(pergunta: str) -> int:
    while True:
        bruto = input(pergunta).strip()
        try:
            return int(bruto)
        except ValueError:
            print("  -> Digite um numero inteiro valido.")


# ---------------------------------------------------------------------------
# Casos de uso do Concierge MVP
# ---------------------------------------------------------------------------

def cadastrar_localidade() -> None:
    print("\n--- Cadastrar localidade ---")
    wb = abrir_planilha()
    ws = wb["Localidades"]
    nome = pedir_texto("Nome da localidade: ")
    media = pedir_numero("Media historica de chuva no mes (mm), se ja souber (0 se nao souber ainda): ")
    ws.append([nome, media])
    salvar_planilha(wb)
    print(f"Localidade '{nome}' cadastrada.")


def cadastrar_cultura() -> None:
    print("\n--- Cadastrar cultura ---")
    wb = abrir_planilha()
    ws = wb["Culturas"]
    nome = pedir_texto("Nome da cultura (ex.: feijao, milho): ")
    limiar = pedir_numero("Limiar minimo de chuva no mes para plantio seguro (mm): ")
    ws.append([nome, limiar])
    salvar_planilha(wb)
    print(f"Cultura '{nome}' cadastrada com limiar de {limiar} mm.")


def cadastrar_familia() -> None:
    print("\n--- Cadastrar familia piloto ---")
    wb = abrir_planilha()
    ws = wb["Familias"]
    novo_id = proximo_id(ws)
    nome = pedir_texto("Nome da familia / responsavel: ")
    localidade = pedir_texto("Localidade (deve existir em 'Localidades'): ")
    cultura = pedir_texto("Cultura principal (deve existir em 'Culturas'): ")
    contato = pedir_texto("Contato (telefone/WhatsApp): ")
    ws.append([novo_id, nome, localidade, cultura, contato])
    salvar_planilha(wb)
    print(f"Familia '{nome}' cadastrada com id {novo_id}.")


def registrar_chuva() -> None:
    print("\n--- Registrar chuva acumulada do mes (dado publico INMET/INPE) ---")
    wb = abrir_planilha()
    ws = wb["RegistroChuva"]
    novo_id = proximo_id(ws)
    localidade = pedir_texto("Localidade: ")
    mes_referencia = pedir_texto("Mes de referencia (ex.: 2026-09): ")
    volume = pedir_numero("Volume de chuva acumulado no mes (mm): ")
    ws.append([novo_id, localidade, mes_referencia, volume, date.today().isoformat()])
    salvar_planilha(wb)
    print(f"Registro de chuva salvo para '{localidade}' em {mes_referencia}: {volume} mm.")


def registrar_nivel_reservatorio() -> None:
    print("\n--- Registrar nivel do reservatorio (relato semanal da familia) ---")
    wb = abrir_planilha()
    familias = linhas_como_dicts(wb["Familias"])
    if not familias:
        print("Nenhuma familia cadastrada ainda. Cadastre uma familia primeiro.")
        return

    print("Familias cadastradas:")
    for f in familias:
        print(f"  [{f['id']}] {f['nome']} - {f['localidade']}")

    familia_id = pedir_inteiro("Id da familia: ")
    if not any(f["id"] == familia_id for f in familias):
        print("Id de familia nao encontrado.")
        return

    nivel = pedir_numero("Nivel do reservatorio relatado (%, de 0 a 100): ")
    observacao = input("Observacao (opcional, Enter para pular): ").strip()

    ws = wb["RegistroReservatorio"]
    novo_id = proximo_id(ws)
    ws.append([novo_id, familia_id, date.today().isoformat(), nivel, observacao])
    salvar_planilha(wb)
    print("Nivel registrado com sucesso.")

    if nivel <= 25:
        print(">> Atencao: nivel critico (<=25%). Considere visitar essa familia com prioridade.")


def _chuva_mais_recente(wb: Workbook, localidade: str) -> dict | None:
    registros = [r for r in linhas_como_dicts(wb["RegistroChuva"]) if r["localidade"] == localidade]
    if not registros:
        return None
    registros.sort(key=lambda r: r["mes_referencia"])
    return registros[-1]


def gerar_sugestao_plantio() -> None:
    print("\n--- Gerar sugestao consultiva de plantio ---")
    wb = abrir_planilha()
    familias = linhas_como_dicts(wb["Familias"])
    if not familias:
        print("Nenhuma familia cadastrada ainda.")
        return

    print("Familias cadastradas:")
    for f in familias:
        print(f"  [{f['id']}] {f['nome']} - {f['localidade']} - cultura: {f['cultura_principal']}")

    familia_id = pedir_inteiro("Id da familia: ")
    familia = next((f for f in familias if f["id"] == familia_id), None)
    if familia is None:
        print("Id de familia nao encontrado.")
        return

    chuva = _chuva_mais_recente(wb, familia["localidade"])
    if chuva is None:
        print(f"Nao ha registro de chuva para a localidade '{familia['localidade']}'. "
              f"Registre a chuva do mes antes de gerar a sugestao.")
        return

    culturas = linhas_como_dicts(wb["Culturas"])
    cultura = next((c for c in culturas if c["nome"] == familia["cultura_principal"]), None)
    if cultura is None:
        print(f"A cultura '{familia['cultura_principal']}' nao esta cadastrada em 'Culturas'.")
        return

    localidades = linhas_como_dicts(wb["Localidades"])
    localidade_info = next((l for l in localidades if l["nome"] == familia["localidade"]), None)
    media_historica = localidade_info["media_historica_mm"] if localidade_info else 0

    chuva_acumulada = chuva["volume_mm"]
    limiar = cultura["limiar_chuva_mm"]
    percentual_media = (chuva_acumulada / media_historica * 100) if media_historica else None

    # --- Logica de sugestao: sempre consultiva, nunca uma determinacao (RN12) ---
    razao = (
        f"chuva acumulada de {chuva_acumulada:.1f} mm em {familia['localidade']} "
        f"({chuva['mes_referencia']}), contra um limiar de referencia de {limiar:.1f} mm "
        f"para {cultura['nome']}"
    )
    if percentual_media is not None:
        razao += f", o que representa {percentual_media:.0f}% da media historica da localidade"

    if chuva_acumulada >= limiar:
        tom = (
            f"O dado disponivel ({razao}) indica uma condicao favoravel para o plantio de "
            f"{cultura['nome']}. Isso e uma leitura com base em chuva e historico, nao uma "
            f"garantia — combine com o que voce esta observando na sua propria terra antes de decidir."
        )
    elif chuva_acumulada >= 0.7 * limiar:
        tom = (
            f"O dado disponivel ({razao}) sugere cautela: a chuva esta abaixo do ideal, mas nao "
            f"criticamente baixa. Pode valer a pena esperar mais alguns dias de chuva antes de "
            f"decidir, mas a palavra final sobre o momento certo e sua."
        )
    else:
        tom = (
            f"O dado disponivel ({razao}) aponta um cenario de chuva bem abaixo do esperado para "
            f"{cultura['nome']}. Vale considerar adiar o plantio ou reduzir a area plantada este mes "
            f"— mas essa e uma sugestao baseada em dado, nao uma determinacao: voce conhece sua terra "
            f"melhor do que qualquer sistema."
        )

    ws = wb["SugestoesPlantio"]
    novo_id = proximo_id(ws)
    ws.append([
        novo_id, familia_id, date.today().isoformat(), cultura["nome"],
        chuva_acumulada, media_historica, percentual_media, tom,
    ])
    salvar_planilha(wb)

    print("\n" + "=" * 70)
    print(f"Sugestao para {familia['nome']} (enviar por voz/WhatsApp/visita):")
    print("-" * 70)
    print(tom)
    print("=" * 70)


def ver_painel_comunitario() -> None:
    """Lista o ultimo nivel registrado de cada familia, do mais critico ao mais tranquilo."""
    print("\n--- Painel comunitario (ultimo nivel de cada familia) ---")
    wb = abrir_planilha()
    familias = {f["id"]: f for f in linhas_como_dicts(wb["Familias"])}
    registros = linhas_como_dicts(wb["RegistroReservatorio"])

    if not registros:
        print("Ainda nao ha nenhum registro de nivel de reservatorio.")
        return

    ultimo_por_familia: dict[int, dict] = {}
    for r in registros:
        atual = ultimo_por_familia.get(r["familia_id"])
        if atual is None or r["data"] >= atual["data"]:
            ultimo_por_familia[r["familia_id"]] = r

    linhas = []
    for familia_id, registro in ultimo_por_familia.items():
        familia = familias.get(familia_id, {})
        linhas.append((
            registro["nivel_percentual"],
            familia.get("nome", f"id {familia_id}"),
            familia.get("localidade", "-"),
            registro["data"],
        ))

    linhas.sort(key=lambda x: x[0])  # mais critico (nivel mais baixo) primeiro

    print(f"{'Nivel %':>8}  {'Familia':<25} {'Localidade':<20} {'Data'}")
    print("-" * 70)
    for nivel, nome, localidade, data_reg in linhas:
        alerta = " <<< CRITICO" if nivel <= 25 else ""
        print(f"{nivel:>7.0f}%  {nome:<25} {localidade:<20} {data_reg}{alerta}")


# ---------------------------------------------------------------------------
# Menu principal
# ---------------------------------------------------------------------------

MENU = """
==========================================================
 MANEJO CAATINGUEIRO - Ferramenta de Coleta (Concierge MVP)
==========================================================
 1. Cadastrar localidade
 2. Cadastrar cultura
 3. Cadastrar familia piloto
 4. Registrar chuva do mes (dado publico INMET/INPE)
 5. Registrar nivel do reservatorio (relato semanal)
 6. Gerar sugestao consultiva de plantio
 7. Ver painel comunitario (niveis, do mais critico ao mais tranquilo)
 0. Sair
----------------------------------------------------------
"""


def main() -> None:
    inicializar_planilha()
    acoes = {
        "1": cadastrar_localidade,
        "2": cadastrar_cultura,
        "3": cadastrar_familia,
        "4": registrar_chuva,
        "5": registrar_nivel_reservatorio,
        "6": gerar_sugestao_plantio,
        "7": ver_painel_comunitario,
    }
    while True:
        print(MENU)
        escolha = input("Escolha uma opcao: ").strip()
        if escolha == "0":
            print("Ate a proxima. Dados salvos em:", ARQUIVO_DADOS)
            break
        acao = acoes.get(escolha)
        if acao is None:
            print("Opcao invalida.")
            continue
        try:
            acao()
        except Exception as erro:  # nao deixa o piloto cair por um erro de digitacao
            print(f"Ocorreu um erro nessa operacao: {erro}")


if __name__ == "__main__":
    main()
